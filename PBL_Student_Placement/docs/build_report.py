"""Builds the PBL project report (DOCX) from the code, requirements, Jira backlog, test results and evidence.

    python docs/build_report.py   ->  docs/PBL_Report_Student_Placement_Vineet_Shah.docx

Jira screenshots in evidence/jira were taken on 26-10-2026 in the USPS space.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "jira"))
import backlog  # noqa: E402

OUT = ROOT / "docs" / "PBL_Report_Student_Placement_Vineet_Shah.docx"
BRAND = RGBColor(0x1E, 0x3A, 0x8A)
PASS = RGBColor(0x16, 0x65, 0x34)
FAIL = RGBColor(0xA3, 0x1F, 0x1F)
HEADER_FILL = "DCE4F7"
CODE_FILL = "F4F6F9"
KEYS = json.loads((ROOT / "jira" / "jira_keys.json").read_text(encoding="utf-8"))["keys"]

JIRA_SHOTS = [
    ("01_backlog_sprint3.png", "Backlog - active USPS Sprint 3 with goal, points and assignees"),
    ("02_board_sprint3.png", "Board of USPS Sprint 3 - report task In Progress, the rest Done"),
    ("03_timeline_epics.png", "Timeline - progress of the four epics"),
    ("08_reports_dashboard.png", "Reports - 49 work items by status, type and assignee"),
    ("06_access_members.png", "Space access - administrator and the invited member"),
    ("07_assigned_to_second_member.png", "The 14 work items assigned to the second member (USPS-26 still In Progress)"),
]
BUG_SHOTS = [
    ("04_bug_USPS-46_top.png", "USPS-46 (BUG-01): description, found by, severity, steps, assignee, labels"),
    ("05_bug_USPS-46_resolution.png", "USPS-46: expected / actual, resolution, retest and evidence"),
]
JIRA_SPRINTS = [  # sprint, key, dates, state at the time of the report
    ("USPS Sprint 1", "S1", "26-10-2026 10:30 - 10:50", "Completed"),
    ("USPS Sprint 2", "S2", "26-10-2026 10:50 - 11:05", "Completed"),
    ("USPS Sprint 3", "S3", "26-10-2026 11:05 - 02-11-2026", "Active"),
]
# bug, test case, run that showed the failure, run after the fix
RETESTS = [("BUG-01", "TC_UI_07", "run1", "run2"), ("BUG-01", "TC_UI_11", "run1", "run2"),
           ("BUG-01", "TC_UI_15", "run1", "run2"), ("BUG-01", "TC_UI_37", "run1", "run2"),
           ("BUG-02", "TC_INT_10", "run3", "run4"), ("BUG-03", "TC_INT_11", "run3", "run4"),
           ("BUG-04", "TC_M1_60", "run3", "run4")]


# ---------------- docx helpers ----------------

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def row_flag(row, tag):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement(tag)
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_cell(cell, text, size=9, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = color
        if i < len(lines) - 1:
            run.add_break()


def table(doc, headers, rows, widths, size=9, keep_together=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, size=size, bold=True)
        shade(t.rows[0].cells[i], HEADER_FILL)
    row_flag(t.rows[0], "w:tblHeader")
    for r in rows:
        row = t.add_row()
        row_flag(row, "w:cantSplit")
        for i, v in enumerate(r):
            if isinstance(v, tuple):
                set_cell(row.cells[i], v[0], size=size, bold=True, color=v[1])
            else:
                set_cell(row.cells[i], v, size=size)
    for n, row in enumerate(t.rows):
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        if keep_together and n < len(t.rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    return t


def code(doc, text, size=8, numbered=True, width=17.0):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Cm(width)
    shade(cell, CODE_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        if numbered:
            r = p.add_run("%3d  " % (i + 1))
            r.font.name = "Consolas"
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0x8A, 0x93, 0xA3)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(size)
        if i < len(lines) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = BRAND
    return h


def para(doc, text, bold_prefix=None, size=11, after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(size)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    return p


def bullet(doc, text, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(size)


class Figures:
    def __init__(self):
        self.n = 0

    def add(self, doc, path, caption, width=None, height=None):
        self.n += 1
        doc.add_picture(str(path), width=Cm(width) if width else None, height=Cm(height) if height else None)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run("Figure %d: %s" % (self.n, caption))
        r.italic = True
        r.font.size = Pt(9.5)

    def pair(self, doc, items, width=7.9):
        """Two screenshots side by side, each with its own caption."""
        t = doc.add_table(rows=2, cols=len(items))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        row_flag(t.rows[0], "w:cantSplit")
        for i, (path, caption) in enumerate(items):
            self.n += 1
            pic = t.rows[0].cells[i].paragraphs[0]
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(str(path), width=Cm(width))
            cap = t.rows[1].cells[i].paragraphs[0]
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run("Figure %d: %s" % (self.n, caption))
            r.italic = True
            r.font.size = Pt(9)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def pairs(self, doc, folder, shots, width=7.9):
        found = [(folder / f, c) for f, c in shots if (folder / f).exists()]
        for i in range(0, len(found), 2):
            self.pair(doc, found[i:i + 2], width)
        return len(found)


def section(doc, landscape):
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    a, b = Cm(29.7), Cm(21.0)
    sec.page_width, sec.page_height = (a, b) if landscape else (b, a)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6 if landscape else 2.0))


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def read(rel):
    return (ROOT / rel).read_text(encoding="utf-8")


def line_count(rel):
    return len(read(rel).rstrip("\n").split("\n"))


def requirements(prefix):
    """(id, text) rows of docs/requirements.md for one module, without markdown markup."""
    rows = []
    for line in read("docs/requirements.md").split("\n"):
        m = re.match(r"\| (FR-%s\.\d+) \| (.*) \|$" % prefix, line)
        if m:
            rows.append([m.group(1), m.group(2).replace("**", "").replace("`", "")])
    return rows


def source_block(rel, start_marker, end_marker):
    src = read(rel).split("\n")
    start = next(i for i, l in enumerate(src) if start_marker in l)
    end = next(i for i, l in enumerate(src) if i > start and end_marker in l)
    return "\n".join(src[start:end]).rstrip()


def counts(run):
    rs = run["results"]
    return len(rs), sum(r["status"] == "PASS" for r in rs), sum(r["status"] == "FAIL" for r in rs)


# ---------------- content ----------------

def main():
    runs = {k: json.loads(read("results/%s.json" % k)) for k in ("run1", "run2", "run3", "run4")}
    by_run = {k: {r["id"]: r for r in v["results"]} for k, v in runs.items()}
    bugs = json.loads(read("docs/bugs.json"))
    results = runs["run4"]["results"]
    figs = Figures()

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))

    # ----- title page -----
    for text, size, bold, color, before in (
            ("Software Engineering Lab - Project-Based Learning (PBL)", 13, False, None, 40),
            ("Unified Student Placement System (USPS)", 24, True, BRAND, 30),
            ("Job Posting & Eligibility Verification  |  Job Application & Tracking", 13, False, None, 6),
            ("Project Report", 13, True, None, 30)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    table(doc, ["Particular", "Details"], [
        ["Name", "Vineet Shah"],
        ["Class / Batch", "5S CSE / A6 - B2"],
        ["Roll No.", "28"],
        ["Course", "Software Engineering Lab (24CS01PR0502)"],
        ["Faculty", "Prof. D. A. Borikar"],
        ["Case study", "Unified System for Student Placement in Academic Institutions (IEEE SRS, Expt-1 to Expt-5)"],
        ["Implemented modules", "Module 1: Job Posting and Eligibility Verification (SRS 4.3, 4.4)\n"
                                "Module 2: Job Application and Tracking (SRS 4.5)"],
        ["Project management", "Jira space USPS - epics, user stories, tasks, sub-tasks, sprints, bugs"],
        ["Date", "26-10-2026"],
    ], [5.0, 12.0], size=11)
    page_break(doc)

    # ----- 1. problem statement -----
    heading(doc, "1. Project Title and Problem Statement")
    para(doc, "Unified Student Placement System (USPS)", bold_prefix="Title: ")
    para(doc, "Campus placement drives are run today with spreadsheets, e-mail threads and notice-board "
              "announcements. Placement officers check every student's eligibility for every company by hand, "
              "students apply to jobs they are not eligible for or miss deadlines, the same data is copied many "
              "times and goes stale, and nobody has a live view of where an application stands.")
    para(doc, "USPS replaces this with one web platform: recruiters post jobs with eligibility criteria, the "
              "placement cell approves them, the system checks every student against the criteria automatically "
              "and with reasons, eligible students apply with a resume before the deadline, and every application "
              "moves through review, shortlisting, interview and offer with a timeline and notifications that "
              "everyone can see in real time.")
    para(doc, "This report covers the two modules implemented for the PBL out of the full system specified in my SE "
              "lab SRS: Module 1 - Job Posting and Eligibility Verification, and Module 2 - Job Application and "
              "Tracking. Registration and login, interview scheduling, offer documents and reports are other "
              "modules of the SRS and are outside this PBL.", italic=True, size=10.5)

    # ----- 2. system / modules -----
    heading(doc, "2. System and Module Description")
    figs.add(doc, ROOT / "docs/figures/architecture.png",
             "Three-tier architecture of the prototype (SRS 2.1); the two implemented modules are shaded", width=16.5)
    table(doc, ["Module", "What it does", "SRS"], [
        ["Module 1 - Job Posting and Eligibility Verification",
         "Recruiter posts a job with package, deadline and an eligibility rule-set; placement officer approves, "
         "rejects or returns it; returned postings are fixed and resubmitted; students search published jobs; "
         "every student is checked against each job automatically with the list of unmet criteria.",
         "4.3, 4.4; FR-09 to FR-14; UC-2; BR-02, BR-03"],
        ["Module 2 - Job Application and Tracking",
         "Eligible student applies with an active resume and cover note and gets a tracking id; recruiter / "
         "officer move the application through the Expt-5 statechart; student withdraws, accepts or declines "
         "an offer; timeline, notifications and live page updates.",
         "4.5; FR-15 to FR-17, FR-25; UC-1; BR-01, BR-04, BR-05, BR-07, BR-08"],
    ], [4.2, 9.6, 3.2], size=9.5)
    table(doc, ["Earlier SE lab work", "Used in the implementation as"], [
        ["Expt-1/2: IEEE SRS, use cases UC-1 and UC-2",
         "Requirements FR-1.x / FR-2.x (docs/requirements.md) with the exception flows E1-E4 and validation rules"],
        ["Expt-3: class diagram",
         "EligibilityCriteria.evaluate() / getReasons(), Application.updateStatus() -> next_state(), "
         "tracking id; tables follow the classes Student, Company, Job, Resume, Application, Notification"],
        ["Expt-4: sequence diagram", "Apply flow: select job -> verify eligibility -> store application -> notify"],
        ["Expt-5: statechart and activity diagram",
         "APPLICATION_TRANSITIONS in app/rules.py is the statechart transition table; the activity "
         "diagram's eligible / not eligible branch is the apply check"],
        ["SRS 6.1 database design", "SQLite tables in app/db.py"],
    ], [5.0, 12.0], size=9.5)
    table(doc, ["Item", "Technology"], [
        ["Language / runtime", "Python 3.11, standard library only (http.server, sqlite3, json, threading, unittest)"],
        ["Database", "SQLite - 10 tables (students, companies, recruiters, placement_officers, resumes, jobs, "
                     "applications, application_history, eligibility_checks, notifications)"],
        ["User interface", "HTML, CSS, JavaScript (fetch API); one page with Student / Recruiter / Placement Officer "
                           "views; polls /api/version every 2 s for live updates"],
        ["Testing", "Python test runner + unittest; user-interface tests in headless Microsoft Edge"],
        ["Project management", "Jira Cloud (team-managed Scrum space USPS)"],
    ], [4.0, 13.0], size=9.5)

    # ----- 3-5. Jira -----
    heading(doc, "3. Jira Project")
    type_counts = {}
    for r in backlog.rows():
        type_counts[r["Issue Type"]] = type_counts.get(r["Issue Type"], 0) + 1
    table(doc, ["Item", "Details"], [
        ["Jira space", "Unified Student Placement System - key USPS (Scrum, team-managed), board USPS board"],
        ["Site", "rbunagpur-team-zhx8jd75.atlassian.net"],
        ["Work items", "%d epics, %d user stories, %d tasks, %d sub-tasks, %d bugs = %d (USPS-1 to USPS-49)" % (
            type_counts["Epic"], type_counts["Story"], type_counts["Task"], type_counts["Subtask"],
            type_counts["Bug"], sum(type_counts.values()))],
        ["Hierarchy", "Epic -> Story / Task / Bug -> Sub-task; USPS-1 to USPS-4 are the epics"],
        ["Workflow", "To Do -> In Progress -> In Review -> Done"],
        ["Members", "A6_B2_28_Vineet Shah (24shahv@rbunagpur.in) - administrator; Vineet Shah (my second, "
                    "personal account, invited by e-mail) - member with 14 work items assigned, including the open USPS-26"],
        ["Status on 26-10-2026", "47 Done; 2 In Progress (USPS-26 report and presentation, and its epic USPS-4)"],
    ], [4.0, 13.0], size=9.5)
    para(doc, "The backlog is kept in jira/backlog.py (also exported as jira/jira_import.csv). The work items were "
              "created in the USPS space with parent links, story points, priorities and sprints, then moved through "
              "the workflow as the work was done. Each bug went To Do -> In Progress -> In Review -> Done and has a "
              "fix-and-retest comment. Issue keys are listed in jira/jira_keys.json.", size=10)
    table(doc, ["Epic", "Jira key", "Description", "Label"],
          [[e[1], KEYS[e[0]], e[2], e[3]] for e in backlog.EPICS], [4.4, 1.8, 9.0, 1.8], size=9)
    sprint_rows = []
    for name, key, dates, state in JIRA_SPRINTS:
        items = [i for i in backlog.ITEMS if i[7] == key]
        keys = ", ".join(KEYS[i[0]] for i in items)
        if key == "S3":
            keys += ", %s to %s (bugs)" % (KEYS["BUG-01"], KEYS["BUG-04"])
        sprint_rows.append([name + "\n" + backlog.SPRINT_GOALS[key], dates, keys,
                            str(sum(i[6] for i in items)), state])
    table(doc, ["Sprint and goal", "Dates", "Work items", "Points", "State"], sprint_rows,
          [5.4, 3.0, 5.6, 1.4, 1.6], size=9)
    figs.pairs(doc, ROOT / "evidence/jira", JIRA_SHOTS)

    heading(doc, "4. Product Backlog")
    assigned = {"US4", "US8", "US10", "US11", "T5", "T7", "T9", "T11"}
    table(doc, ["Jira key", "Type", "Summary", "Epic", "Priority", "Points", "Sprint", "Assignee"],
          [[KEYS[i[0]], i[2], i[3], KEYS[i[1]], i[5], str(i[6]), i[7], "VS" if i[0] in assigned else "AS"]
           for i in backlog.ITEMS],
          [1.7, 1.3, 6.4, 1.5, 1.6, 1.2, 1.2, 1.6], size=8.5, keep_together=False)
    para(doc, "AS = A6_B2_28_Vineet Shah, VS = Vineet Shah (second account). Sub-tasks (%d): %s." % (
        len(backlog.SUBTASKS), "; ".join("%s %s (%s)" % (KEYS["ST%d" % n], s, KEYS[p])
                                         for n, (p, s) in enumerate(backlog.SUBTASKS, 1))), size=9)

    heading(doc, "5. User Stories")
    table(doc, ["Key", "User story and acceptance criteria", "Points"],
          [[KEYS[i[0]], i[3] + "\n" + i[4], str(i[6])] for i in backlog.ITEMS if i[2] == "Story"],
          [1.8, 13.6, 1.6], size=9, keep_together=False)

    # ----- 6. implementation -----
    heading(doc, "6. Implementation of the Two Modules")
    heading(doc, "6.1 Module 1: Job Posting and Eligibility Verification", 2)
    para(doc, "The rules of the module are in app/rules.py and have no HTTP or SQL code, so they are unit-tested "
              "directly. validate_job_posting() checks a posting field by field and returns the first problem; "
              "EligibilityCriteria (class of the Expt-3 diagram) evaluates a student and returns every unmet "
              "criterion as a reason. app/server.py adds the actor checks, stores the job as Pending, handles the "
              "officer's decision and notifies the officers, the recruiter and the eligible students.")
    table(doc, ["ID", "Functional requirement (Module 1)"], requirements("1"), [1.6, 15.4], size=8.5,
          keep_together=False)
    para(doc, "EligibilityCriteria - automatic eligibility check with reasons (FR-1.7):", size=10)
    code(doc, source_block("app/rules.py", "class EligibilityCriteria:", "def to_dict(self):"), size=7.8,
         numbered=False)

    heading(doc, "6.2 Module 2: Job Application and Tracking", 2)
    para(doc, "An application follows the statechart drawn in Expt-5. The statechart is coded as a transition "
              "table (APPLICATION_TRANSITIONS) and every status change goes through next_state(), so a step that the "
              "diagram does not allow (for example an offer before selection, BR-08) is refused. apply() makes the "
              "UC-1 checks in order and saves the application, its history and the notifications in one "
              "transaction; state-changing requests run under one lock so a double click is applied once.")
    figs.add(doc, ROOT / "docs/figures/statechart_expt5.png",
             "Application statechart from my Expt-5 submission - implemented as APPLICATION_TRANSITIONS", width=16.0)
    table(doc, ["ID", "Functional requirement (Module 2)"], requirements("2"), [1.6, 15.4], size=8.5,
          keep_together=False)
    para(doc, "PlacementApp.apply() - UC-1 checks, tracking id and rollback (FR-2.1 to FR-2.4):", size=10)
    code(doc, source_block("app/server.py", "    def apply(self, body):", "    def _history(self"), size=7.3,
         numbered=False)

    heading(doc, "6.3 REST API, database and pages", 2)
    table(doc, ["Method and path", "Purpose"], [
        ["GET /api/meta, /api/version", "Demo users, resumes, statechart; data version polled by every page"],
        ["POST /api/jobs, PUT /api/jobs/<id>", "Post a job / resubmit a returned job (Module 1)"],
        ["POST /api/jobs/<id>/decision", "Officer approves, rejects or returns a posting (Module 1)"],
        ["GET /api/jobs?q=&location=&min_ctc=&eligible_only=&page=", "Job board with eligibility (Module 1)"],
        ["GET /api/jobs/<id>/eligibility?student_id=", "Eligibility check with reasons, recorded (Module 1)"],
        ["GET /api/jobs/manage?recruiter_id= | status=", "Recruiter's postings / officer's pending list"],
        ["POST /api/applications", "Apply (Module 2)"],
        ["POST /api/applications/<id>/events, /withdraw", "Statechart step, offer response, withdrawal (Module 2)"],
        ["GET /api/applications?student_id= | recruiter_id=, /api/applications/<id>",
         "My applications / applicants of own jobs (BR-07) / timeline"],
        ["GET /api/notifications?recipient_type=&recipient_id=", "Notifications, newest first"],
    ], [7.0, 10.0], size=8.5)
    figs.pair(doc, [(ROOT / "evidence/ui/TC_UI_03.png", "Student 105: unmet criteria listed, Apply disabled"),
                    (ROOT / "evidence/ui/TC_UI_26.png", "Recruiter posts J-5; it is stored as Pending")])
    figs.pair(doc, [(ROOT / "evidence/ui/TC_UI_31.png", "Officer returns J-3 with remarks"),
                    (ROOT / "evidence/ui/TC_UI_35.png", "Recruiter moves USPS-26-0001 to Offer Issued")])
    figs.pair(doc, [(ROOT / "evidence/ui/TC_UI_38.png", "Student timeline after accepting the offer"),
                    (ROOT / "evidence/ui/TC_UI_39.png", "Live update: a job posted elsewhere appears on the officer's "
                                                        "open page with a pop-up")])
    para(doc, "Integration: a posting made on the recruiter page goes through the API into Module 1, is stored as "
              "Pending and appears on the officer page within 2 seconds; after approval Module 1 evaluates every "
              "student and notifies the eligible ones, whose job boards update by themselves; applying runs the "
              "Module 1 eligibility check again inside Module 2 before the application is stored.", size=10.5)

    # ----- 7. source code -----
    heading(doc, "7. Source Code")
    files = [("app/rules.py", "Business rules of both modules: posting validation, EligibilityCriteria, decisions, "
                              "statechart, tracking id"),
             ("app/server.py", "REST API, actor checks, Module 1 and Module 2 services, live-update version"),
             ("app/db.py", "SQLite schema (SRS 6.1), demo data, transactions"),
             ("app/static/index.html", "Page with student, recruiter and officer views"),
             ("app/static/app.js", "Page logic: forms, tables, timeline, live polling"),
             ("app/static/style.css", "Styles, responsive layout"),
             ("run.py", "Starts the portal"),
             ("tests/helpers.py", "Test clock, throw-away server, HTTP client"),
             ("tests/m1_cases.py", "61 Module 1 cases"), ("tests/m2_cases.py", "42 Module 2 cases"),
             ("tests/integration_cases.py", "11 integration cases"), ("tests/ui/ui_cases.js", "41 UI cases"),
             ("tests/run_tests.py", "Executes all cases and records actual results"),
             ("tests/capture_evidence.py", "Screenshots of UI cases for the report"),
             ("jira/backlog.py", "Jira backlog and CSV export")]
    table(doc, ["File", "Lines", "Contents"], [[f, str(line_count(f)), d] for f, d in files], [5.2, 1.4, 10.4],
          size=9)
    code(doc, "python run.py                               # start the portal: http://127.0.0.1:8091\n"
              "python -m tests.run_tests run4              # execute all 155 test cases\n"
              "python -m unittest -v tests.test_all        # Python cases through unittest",
         size=8, numbered=False)

    # ----- 8. test cases (landscape) -----
    section(doc, landscape=True)
    heading(doc, "8. Test Cases and Test Results")
    table(doc, ["Technique", "How it was applied"], [
        ["Unit testing", "app/rules.py tested directly: posting validation, eligibility, decisions, statechart "
                         "(TC_M1_01-37, TC_M2_01-08)"],
        ["Functional testing", "Every functional requirement checked through the API and the page"],
        ["Integration testing", "Module 1 + Module 2 + database + notifications: full placement flow, rollback, "
                                "concurrency, restart, live updates (TC_INT_01-11, TC_UI_39-40)"],
        ["Boundary Value Analysis", "title 2/3/120/121, package 0/0.1/100/100.5, CGPA 6.99/7.00/10/10.01, backlogs, "
                                    "remarks 4/5/201, cover note 500/501, deadline = now / now+1 min, apply 1 min before "
                                    "and at the deadline, page 0"],
        ["Equivalence Partitioning", "valid and invalid classes of every input: numbers as words, unknown branch, "
                                     "wrong batch, inactive / other student's resume, unknown decision or event"],
        ["Validation testing", "every error message of the API and of the page forms"],
    ], [5.0, 21.5], size=9.5)
    para(doc, "Test cases were designed from docs/requirements.md before execution and are automated: Python cases "
              "call app/rules.py or a fresh server with the demo data and a fixed test clock (26-10-2026 10:00); UI "
              "cases fill the page in Microsoft Edge and read what it shows. Actual results below are from the final "
              "run (run4, %s). Earlier runs are compared in Section 10." % runs["run4"]["executed_at"].replace("T", " "),
         size=10)
    by_module = {}
    for r in results:
        by_module.setdefault(r["module"], []).append(r)
    widths = [1.9, 2.6, 4.4, 5.2, 5.3, 5.3, 1.4]
    for module, rs in by_module.items():
        heading(doc, module, level=2)
        rows = [[r["id"], r["module"].split(" (")[0], "%s\n[%s]" % (r["scenario"], r["technique"]), r["data"],
                 r["expected"], r["actual"], (r["status"], PASS if r["status"] == "PASS" else FAIL)] for r in rs]
        table(doc, ["Test Case ID", "Module", "Test Scenario", "Test Data", "Expected Result", "Actual Result",
                    "Status"], rows, widths, size=7.8, keep_together=False)
    heading(doc, "Test execution summary", level=2)
    summary = []
    for module, rs in by_module.items():
        row = [module, str(len(rs))]
        for k in ("run1", "run2", "run3", "run4"):
            got = [by_run[k][r["id"]] for r in rs if r["id"] in by_run[k]]
            row.append("%d / %d" % (sum(x["status"] == "PASS" for x in got), len(got)))
        summary.append(row)
    total = [("Total", None), (str(len(results)), None)]
    for k in ("run1", "run2", "run3", "run4"):
        n, p, _ = counts(runs[k])
        total.append(("%d / %d" % (p, n), None))
    summary.append(total)
    table(doc, ["Module", "Cases", "Run 1 passed", "Run 2 passed (retest)", "Run 3 passed (+4 new cases)",
                "Run 4 passed (retest)"], summary, [7.5, 2.0, 3.8, 4.2, 4.6, 4.2], size=10)

    # ----- 9-10. bugs -----
    heading(doc, "9. Bugs Recorded in Jira")
    para(doc, "Run 1 failed 15 UI cases; they all trace to one defect (BUG-01) because the later UI cases build on "
              "the application that could not be submitted. After the retest, an exploratory session (two officers, "
              "double clicks, an approval after the deadline) found BUG-02 to BUG-04; new test cases TC_INT_10, "
              "TC_INT_11, TC_M1_60 and TC_M1_61 were written for them and failed in run 3. All four bugs were "
              "recorded in Jira (USPS Sprint 3), fixed and retested.", size=10)
    table(doc, ["Bug ID", "Bug description", "Steps to reproduce", "Severity / Priority", "Status", "Resolution"],
          [["%s\n(%s)" % (b["id"], KEYS[b["id"]]),
            "%s\n\nExpected: %s\nActual: %s\nFound by: %s" % (b["summary"], b["expected"], b["actual"], b["found_by"]),
            "\n".join("%d. %s" % (i, s) for i, s in enumerate(b["steps"], 1)),
            "%s / %s" % (b["severity"], b["priority"]), b["status"], b["resolution"] + "\nRetest: " + b["retest"]]
           for b in bugs], [1.7, 7.4, 5.8, 2.4, 1.5, 7.6], size=7.8, keep_together=False)
    figs.pairs(doc, ROOT / "evidence/jira", BUG_SHOTS, width=10.5)

    heading(doc, "10. Bug Fixing and Retesting Evidence")
    table(doc, ["Test case", "Bug", "Before the fix - actual", "Status", "After the fix - actual", "Status"],
          [[tc, "%s\n(%s)" % (bug, KEYS[bug]), "%s: %s" % (before, by_run[before][tc]["actual"]),
            (by_run[before][tc]["status"], FAIL), "%s: %s" % (after, by_run[after][tc]["actual"]),
            (by_run[after][tc]["status"], PASS)] for bug, tc, before, after in RETESTS],
          [2.2, 1.9, 9.0, 1.4, 9.0, 1.4], size=8.5)
    section(doc, landscape=False)
    figs.pair(doc, [(ROOT / "evidence/bugs/BUG-01_before_TC_UI_07.png", "BUG-01 before: no resume pre-selected "
                                                                        "(TC_UI_07 FAIL, run 1)"),
                    (ROOT / "evidence/bugs/BUG-01_after_TC_UI_07.png", "BUG-01 after: active resume pre-selected "
                                                                       "(TC_UI_07 PASS)")])
    figs.pair(doc, [(ROOT / "evidence/bugs/BUG-01_before_TC_UI_11.png", "BUG-01 before: submit refused "
                                                                        "(TC_UI_11 FAIL, run 1)"),
                    (ROOT / "evidence/bugs/BUG-01_after_TC_UI_11.png", "BUG-01 after: application submitted "
                                                                       "(TC_UI_11 PASS)")])
    for bug, text in (("BUG-02", "double approval"), ("BUG-03", "double status change"),
                      ("BUG-04", "approval after the deadline")):
        figs.add(doc, ROOT / ("evidence/bugs/%s_before.png" % bug), "%s (%s) before the fix - run 3" % (bug, text),
                 width=15.5)
        figs.add(doc, ROOT / ("evidence/bugs/%s_after.png" % bug), "%s after the fix - run 4" % bug, width=15.5)
    para(doc, "Console summary of the four runs:", size=10)
    code(doc, "\n".join(read("results/%s-console.txt" % k).rstrip().split("\n")[-1] + "   (%s)" % k
                        for k in ("run1", "run2", "run3", "run4")), size=8, numbered=False)

    # ----- 11. demo -----
    heading(doc, "11. Final Working Demonstration")
    for step in ["Start the portal: python run.py and open http://127.0.0.1:8091 in two browser windows.",
                 "Window 1 as Recruiter (Ritika Menon - Nimbus): post a job. Window 2 as Placement Officer: the job "
                 "appears in Pending job approvals within 2 seconds with a pop-up; approve it.",
                 "Switch window 2 to Student Diya Joshi: the new job is on the job board as Eligible and a 'New job "
                 "for you' notification has arrived; apply with the pre-selected resume and get a tracking id.",
                 "Window 1: the applicant appears; move it through Move to review, Shortlist, Schedule interview, "
                 "Mark selected and Issue offer letter while the student's timeline updates by itself.",
                 "Student accepts the offer, becomes Placed, and a further application is refused (BR-05).",
                 "Show an ineligible student (Rohan, ME) with the unmet criteria and the disabled Apply button.",
                 "Run python -m tests.run_tests to execute all 155 test cases (about 1 minute)."]:
        bullet(doc, step, size=10.5)
    figs.pair(doc, [(ROOT / "evidence/ui/TC_UI_40.png", "Approved job appears on the student's open board"),
                    (ROOT / "evidence/ui/TC_UI_41.png", "Same page on a 390 px phone screen")], width=7.0)

    # ----- 12. conclusion -----
    heading(doc, "12. Conclusion and Future Scope")
    n1, p1, f1 = counts(runs["run1"])
    n4, p4, _ = counts(runs["run4"])
    para(doc, "Two modules of the Unified Student Placement System from my SE lab SRS were implemented: job posting "
              "and eligibility verification, and job application and tracking, as one web application with a REST "
              "API, an SQLite database and live-updating pages for students, recruiters and placement officers. The "
              "design follows the earlier lab work - the use cases, the class diagram and the application "
              "statechart. The work was planned and tracked in Jira as 4 epics, 11 user stories, 11 tasks, 19 "
              "sub-tasks and 3 sprints. %d test cases using functional, unit, integration, boundary value, "
              "equivalence partitioning and validation testing were executed: run 1 passed %d of %d, the defects "
              "were logged in Jira, fixed and retested, and the final run passed all %d." % (n4, p1, n1, n4))
    for f in ["Login and role-based access (SRS 4.1) so the acting user comes from the session, not a selector.",
              "Interview scheduling with rounds and slots, and offer letters as documents (SRS 4.6, 4.7).",
              "E-mail notifications and server-sent events instead of polling for live updates.",
              "Reports and analytics of placement statistics for the placement cell."]:
        bullet(doc, f, size=10.5)

    page_break(doc)
    heading(doc, "Appendix A: app/rules.py (business rules of both modules)")
    code(doc, read("app/rules.py"), size=7.3)

    doc.save(OUT)
    print("saved", OUT, "| figures:", figs.n)


if __name__ == "__main__":
    main()
